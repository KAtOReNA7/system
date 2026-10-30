"""Only synthetic workbook rows are used; no local business files are read."""
import contextlib
import hashlib
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts/workspace"))
import prepare_local_snapshot as snapshot  # noqa: E402


def workbook(path, headers, rows, sheet_name="Sheet1", header_row=1):
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = sheet_name
    for _ in range(header_row - 1):
        sheet.append(["合成测试说明"])
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    book.save(path)
    book.close()


class WorkspaceSnapshotTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="workspace-synthetic-")
        self.root = Path(self.temporary.name)
        for relative in [snapshot.PARTITION, snapshot.CHANNEL, snapshot.PRD, snapshot.WORK_PRD, snapshot.CLASS_PRD, snapshot.TAXONOMY, snapshot.FOUNDATION_CONFIRMATION]:
            target = self.root / relative
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes((ROOT / relative).read_bytes())
        confirmation = json.loads((self.root / snapshot.FOUNDATION_CONFIRMATION).read_text(encoding="utf-8"))
        confirmation["scope"]["workCount"] = 1
        (self.root / snapshot.FOUNDATION_CONFIRMATION).write_text(json.dumps(confirmation), encoding="utf-8")
        contract = json.loads((self.root / snapshot.CHANNEL).read_text(encoding="utf-8"))
        contract["channelMaster"]["expectedRawPairCount"] = 2
        contract["channelMaster"]["expectedCanonicalChannelCount"] = 2
        (self.root / snapshot.CHANNEL).write_text(json.dumps(contract, ensure_ascii=False), encoding="utf-8")
        self.rows = [
            ["2025-05", "1", "示例甲渠道", "音频", "100", "测试甲", 100],
            ["2026-04", "1", "示例甲渠道", "音频", "Y100", "测试甲", -20],
            ["2026-05", "1", "示例甲渠道", "音频", "100", "测试甲", 999],
            ["2026-04", "2", "示例乙渠道", "音频", "200", "测试乙", 300],
            ["2026-04", "2", "示例乙渠道", "音频", "300", "测试丙", 0],
            ["2026-04", "3", "未映射测试渠道", "音频", "400", "测试丁", 10],
            ["2026-04", "2", "示例乙渠道", "音频", "100", "测试甲", 25]
        ]
        self.buyout_indices = {3, 6}
        partition = json.loads((self.root / snapshot.PARTITION).read_text(encoding="utf-8"))
        headers = [*snapshot.REAL_BILL_COLUMNS, "类型"]
        for key, name in partition["workbooks"].items():
            selected = [[*row, "买断" if index in self.buyout_indices else "" if key == "totalLedger" else "分成"] for index, row in enumerate(self.rows) if key == "totalLedger" or (index in self.buyout_indices) == (key == "buyout")]
            workbook(self.root / "data/real-bills" / name, headers, selected)
        mapping = self.root / snapshot.MAPPING
        mapping.parent.mkdir(parents=True)
        mapping.write_text(json.dumps({"effective_mapping_snapshot": [{"raw_work_id": "100", "target_standard_work_id": "400", "status": "candidate_not_applied"}]}), encoding="utf-8")
        prefixes = contract["channelMaster"]["columnPrefixes"]
        workbook(self.root / snapshot.CHANNEL_BOOK, list(prefixes.values()), [
            ["1", "示例甲渠道", "规范甲渠道", "终端销售平台", "会员订阅", "有声书/音频", "已确认", ""],
            ["2", "示例乙渠道", "规范乙渠道", "终端销售平台", "会员订阅", "有声书/音频", "已确认", ""]
        ], "渠道补全表", 4)
        workbook(self.root / "data/master-data/数字版权台账.xlsx", ["作品ID", "出版书名", "合同书名", "作者署名", "签订日期", "到期时间", "产品线"], [["100", "不替换账单名称", "另一标题", "测试作者", "2020-01-02", "2030-12-31", "不能推断分类"]])
        workbook(self.root / "data/master-data/原创全库2.xlsx", ["作品ID", "作者笔名", "一级分类", "二级分类", "三级分类", "授权时间", "结束时间"], [["100", "测试作者", "测试一级", "测试二级", "测试三级", "2020-01-02", "2030-12-31"]])
        workbook(self.root / "data/master-data/授权汇总台账.xlsx", ["作品ID", "作者", "出版or原创", "我方版权到期时间"], [["100", "测试作者", "原创", "2030-12-31"]])
        workbook(self.root / "data/master-data/授权关系仪表板.xlsx", ["我方作品ID", "作者名称"], [["100", "测试作者"]])
        foundation_headers = ["序号", "书名", "作品编号", "作者", "一级分类（最终采用）", "二级分类（最终采用）", "三级分类（最终采用）", "辅助标签（最终采用）", "备注"]
        workbook(self.root / snapshot.FOUNDATION, foundation_headers, [[1, "测试甲", "100", "测试作者", "出版物", "小说", "外国小说", "无", ""]], "01_出版物")
        book = load_workbook(self.root / snapshot.FOUNDATION)
        book.create_sheet("02_网文").append(foundation_headers)
        book.save(self.root / snapshot.FOUNDATION)
        book.close()

    def tearDown(self):
        self.temporary.cleanup()

    def prepare(self):
        result = snapshot.prepare(self.root)
        envelope = json.loads((self.root / snapshot.OUTPUT).read_text(encoding="utf-8"))
        payload = json.loads(envelope["payloadJson"])
        self.assertEqual(result["status"], "READY")
        self.assertEqual(hashlib.sha256(envelope["payloadJson"].encode()).hexdigest(), envelope["sha256"])
        return payload

    def test_posting_reversals_missing_months_and_incomplete_month_are_distinct(self):
        payload = self.prepare()
        item = next(row for row in payload["items"] if row["id"] == "100")
        self.assertEqual(item["salesShareCash12Months"], 80)
        self.assertEqual(item["cashCoverage"], {"observedMonths": 2, "expectedMonths": 12})
        self.assertEqual(item["completeBillMonths"], 2)
        self.assertEqual(item["cashType"], "mixed")
        self.assertEqual(item["buyoutAudit"]["cash12Months"], 25)
        history = {row["month"]: row for row in item["monthlyCash"]}
        self.assertIsNone(history["2025-06"]["amount"])
        self.assertEqual(history["2026-04"]["amount"], -20)
        self.assertEqual(history["2026-05"]["amount"], 999)
        self.assertFalse(history["2026-05"]["complete"])
        self.assertEqual(sum(row["amount"] or 0 for row in item["channels"]), 80)
        self.assertTrue(all(row["pointEstimate"] is None for row in item["predictions"]))
        self.assertEqual(payload["dataset"]["billMaxMonth"], "2026-05")

    def test_candidate_mapping_is_not_applied_and_fields_require_exact_identity(self):
        items = {row["id"]: row for row in self.prepare()["items"]}
        self.assertEqual(set(items), {"100", "200", "300", "400"})
        self.assertEqual(items["100"]["rawWorkIds"], ["100", "Y100"])
        self.assertEqual(set(items["100"]["businessForms"]), {"有声版权", "有声成品"})
        self.assertEqual(items["100"]["workName"], "测试甲")
        self.assertEqual(items["100"]["billNames"], ["测试甲"])
        self.assertEqual(items["100"]["authorName"], "测试作者")
        self.assertEqual(items["100"]["categoryPath"], ["出版物", "小说", "外国小说"])
        self.assertEqual(items["100"]["rightsEndDate"], "2030-12-31")
        self.assertTrue(any(row["code"] == "PENDING_MAPPING_NOT_APPLIED" for row in items["100"]["qualityIssues"]))
        self.assertIsNone(items["400"]["authorName"])
        self.assertTrue(any(row["code"] == "CHANNEL_MAPPING_MISSING" for row in items["400"]["qualityIssues"]))
        self.assertEqual(items["400"]["salesShareCash12Months"], 10)

    def test_real_zero_and_pure_buyout_null_are_not_interchanged(self):
        items = {row["id"]: row for row in self.prepare()["items"]}
        self.assertEqual(items["300"]["salesShareCash12Months"], 0)
        self.assertEqual(items["300"]["cashCoverage"]["observedMonths"], 1)
        item = items["200"]
        self.assertIsNone(item["salesShareCash12Months"])
        self.assertEqual(item["buyoutAudit"]["totalCash"], 300)
        self.assertTrue(all(row["amount"] is None for row in item["monthlyCash"]))
        self.assertTrue(all(row["reason"] == "buyout_outside_m2_forecast_scope" for row in item["predictions"]))

    def test_digital_rights_have_priority_over_other_sources_and_internal_conflicts_stay_null(self):
        workbook(self.root / "data/master-data/授权汇总台账.xlsx", ["作品ID", "作者", "我方版权到期时间"], [["100", "另一个测试作者", "2031年"]])
        item = next(row for row in self.prepare()["items"] if row["id"] == "100")
        self.assertIsNone(item["authorName"])
        self.assertEqual(item["rightsEndDate"], "2030-12-31")
        self.assertTrue(any(row["code"] == "MASTER_FIELD_CONFLICT" for row in item["qualityIssues"]))
        self.assertTrue(any(row["code"] == "LOWER_PRIORITY_RIGHTS_DIFFER" for row in item["qualityIssues"]))
        book = load_workbook(self.root / "data/master-data/数字版权台账.xlsx")
        book.active.append(["Y100", "测试甲", "", "测试作者", "2020-01-02", "2032-12-31", ""])
        book.save(self.root / "data/master-data/数字版权台账.xlsx")
        book.close()
        item = next(row for row in self.prepare()["items"] if row["id"] == "100")
        self.assertIsNone(item["rightsEndDate"])

    def test_authority_checks_and_every_read_source_are_bound(self):
        payload = self.prepare()
        self.assertEqual(len(payload["authorityChecks"]), 7)
        self.assertTrue(all(payload["authorityChecks"].values()))
        self.assertEqual(len(payload["sourceDirectories"]), 5)
        required = {"total-ledger-authority", "sales-share-ledger-authority", "buyout-ledger-authority", "user-reviewed-channel-master"}
        self.assertTrue(required.issubset({row["role"] for row in payload["sources"]}))
        for source in payload["sources"]:
            actual = snapshot.source_metadata(self.root, source["path"])
            self.assertEqual({key: source[key] for key in actual}, actual)
            self.assertFalse(Path(source["path"]).is_absolute())

    def test_partition_failure_and_missing_source_preserve_previous_snapshot(self):
        self.prepare()
        target = self.root / snapshot.OUTPUT
        before = target.read_bytes()
        source = self.root / "data/real-bills/渠道实销汇总 -分成.xlsx"
        book = load_workbook(source)
        book.active["G2"] = 123
        book.save(source)
        book.close()
        with self.assertRaises(Exception):
            snapshot.prepare(self.root)
        self.assertEqual(target.read_bytes(), before)
        source.unlink()
        with self.assertRaises(Exception):
            snapshot.prepare(self.root)
        self.assertEqual(target.read_bytes(), before)

    def test_source_change_during_read_is_refused(self):
        self.prepare()
        before = (self.root / snapshot.OUTPUT).read_bytes()
        original = snapshot.read_master_fields
        def changing_reader(root):
            result = original(root)
            (root / "data/master-data/new-authority.xlsx").write_bytes(b"synthetic new file")
            return result
        with patch.object(snapshot, "read_master_fields", changing_reader):
            with self.assertRaisesRegex(snapshot.SnapshotError, "SOURCE_CHANGED_DURING_READ"):
                snapshot.prepare(self.root)
        self.assertEqual((self.root / snapshot.OUTPUT).read_bytes(), before)

    def test_missing_amount_does_not_become_zero(self):
        for value in [None, float("nan"), float("inf"), ""]:
            frame = pd.DataFrame([["2026-04", "1", "示例渠道", "音频", "100", "示例作品", value]], columns=snapshot.REAL_BILL_COLUMNS)
            with self.assertRaisesRegex(snapshot.SnapshotError, "INVALID_BILL_AMOUNT"):
                snapshot.build_items({"salesShare": frame, "buyout": frame.iloc[0:0]}, {}, {}, {}, set(), "2026-04")
        frame = pd.DataFrame([["2026-04", "1", "示例渠道", "音频", "invalid-work-id", "示例作品", 5]], columns=snapshot.REAL_BILL_COLUMNS)
        with self.assertRaisesRegex(snapshot.SnapshotError, "INVALID_WORK_ID"):
            snapshot.build_items({"salesShare": frame, "buyout": frame.iloc[0:0]}, {}, {}, {}, set(), "2026-04")

    def test_explicit_audit_preserves_strict_cache_and_locates_extra_records(self):
        self.prepare()
        strict_before = (self.root / snapshot.OUTPUT).read_bytes()
        source = self.root / "data/real-bills/渠道实销汇总 -分成.xlsx"
        book = load_workbook(source)
        book.active.append(["2026-05", "1", "示例甲渠道", "音频", "100", "测试甲", 17, "分成"])
        extra_row = book.active.max_row
        book.save(source)
        book.close()
        with self.assertRaises(Exception):
            snapshot.prepare(self.root)
        result = snapshot.prepare(self.root, audit_unreconciled=True)
        self.assertEqual(result["status"], "AUDIT_READY")
        self.assertEqual((self.root / snapshot.OUTPUT).read_bytes(), strict_before)
        envelope = json.loads((self.root / snapshot.AUDIT_OUTPUT).read_text(encoding="utf-8"))
        self.assertEqual(envelope["schema"], "workspace.local-audit.v1")
        payload = json.loads(envelope["payloadJson"])
        self.assertEqual(payload["reconciliation"], {"missingRowCount": 0, "extraRowCount": 1, "missingByMonth": {}, "extraByMonth": {"2026-05": 1}})
        self.assertFalse(payload["dataset"]["cashAuthorityVerified"])
        self.assertEqual(payload["dataset"]["kind"], "RAW_LEDGER_AUDIT")
        self.assertFalse(payload["authorityChecks"]["amountConserved"])
        self.assertTrue(all(row["dataStatus"] == "needs_review" for row in payload["items"]))
        item = next(row for row in payload["items"] if row["id"] == "100")
        self.assertEqual(item["salesShareCash12Months"], 80)
        self.assertTrue(any(row["sourceLabel"].endswith(f"第{extra_row}行") for row in item["qualityIssues"]))

    def test_audit_is_not_an_automatic_or_general_error_fallback(self):
        with self.assertRaisesRegex(snapshot.SnapshotError, "AUDIT_REQUIRES_RECONCILIATION_DIFFERENCE"):
            snapshot.prepare(self.root, audit_unreconciled=True)
        source = self.root / "data/real-bills/渠道实销汇总 -分成.xlsx"
        book = load_workbook(source)
        book.active["H2"] = "买断"
        book.save(source)
        book.close()
        with self.assertRaisesRegex(snapshot.SnapshotError, "AUDIT_SPLIT_TYPE_NOT_PURE"):
            snapshot.prepare(self.root, audit_unreconciled=True)
        self.assertFalse((self.root / snapshot.AUDIT_OUTPUT).exists())

    def test_uncontrolled_foundation_classification_cannot_fall_back_to_older_sources(self):
        book = load_workbook(self.root / snapshot.FOUNDATION)
        book["01_出版物"]["E2"] = "非受控一级分类"
        book.save(self.root / snapshot.FOUNDATION)
        book.close()
        item = next(row for row in self.prepare()["items"] if row["id"] == "100")
        self.assertIsNone(item["category"])
        self.assertEqual(item["categoryPath"], [])
        self.assertTrue(any(row["code"] == "CLASSIFICATION_OUTSIDE_CONTROLLED_TREE" for row in item["qualityIssues"]))


if __name__ == "__main__":
    unittest.main(verbosity=2)
